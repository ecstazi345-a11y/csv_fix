"""Read-only workbook of a Constructor result before admission.

Does not write a database, does not approve the month, and does not invent
masters, crews, or missing scope values.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from agents.control_room.dtos import (
    ConstructorDecisionRegistryView,
    ConstructorProfessionalResultView,
    ConstructorRowDecisionView,
    constructor_row_awaits_human_review,
)
from agents.control_room.presentation import constructor_row_outcome_label
from agents.human_surface.mission_draft import WorkingContext
from utils.month_key import normalize_month_key


ADMISSION_STATUS = "NOT_CHECKED"
PRE_ADMISSION_TITLE = "ПРЕДВАРИТЕЛЬНОЕ МЕСЯЧНОЕ НАМЕРЕНИЕ"
PRE_ADMISSION_STATUS = "БЕЗ ПРОВЕРКИ ПРОИЗВОДСТВЕННОЙ ГОТОВНОСТИ И ДОПУСКА"
PRE_ADMISSION_NOTICE = PRE_ADMISSION_STATUS
EXPORT_SECTION_TITLE = "Временный выпуск без допуска"
EXPORT_INTRO_RU = (
    "Предварительный производственный выпуск "
    "без проверки производственной готовности и допуска."
)
EXPORT_UNAVAILABLE_RU = (
    "Выгрузка станет доступна после формирования "
    "предварительного состава Constructor."
)
EXPORT_BUTTON_LABEL = "ВЫГРУЗИТЬ ПРЕДВАРИТЕЛЬНОЕ МЕСЯЧНОЕ НАМЕРЕНИЕ · XLSX"
EXPORT_WARNING_RU = EXPORT_INTRO_RU
EMPLOYEE_ROLE_RU = "Инженер формирования месячного плана"

_FIELD_COLUMNS = (
    "№",
    "Титул",
    "Система",
    "Пакет / IWP",
    "BOQ-код",
    "Наименование работы",
    "Ед. изм.",
    "Остаток",
    "Объём предварительного месячного намерения",
    "Ответственный мастер",
    "Звено / бригада",
    "Неделя 1",
    "Неделя 2",
    "Неделя 3",
    "Неделя 4",
    "Неделя 5",
    "Примечание",
)
_EXCLUDED_COLUMNS = (
    "Титул",
    "Дисциплина",
    "Система",
    "Пакет / IWP",
    "BOQ-код",
    "Наименование работы",
    "Решение Constructor",
    "Причина исключения",
)
_ATTENTION_COLUMNS = (
    "Титул",
    "Дисциплина",
    "Система",
    "Пакет / IWP",
    "BOQ-код",
    "Наименование работы",
    "Доступный объём",
    "Почему требуется решение",
    "Статус",
)
_SUMMARY_COLUMNS = (
    "Дисциплина",
    "Количество рабочих строк",
    "Объём работ",
    "Ед. изм.",
    "Количество титулов",
    "Количество систем",
    "Количество пакетов",
)
_WEEK_COLUMNS = 5
_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_NOTICE_FILL = PatternFill("solid", fgColor="F4B183")
_HEADER_FILL = PatternFill("solid", fgColor="D6DCE4")
_TITLE_FONT = Font(bold=True, color="FFFFFF", size=14)
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_NUMBER_RE = re.compile(r"^-?\d+(?:[.,]\d+)?$")
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def export_available(registry: Optional[ConstructorDecisionRegistryView]) -> bool:
    if registry is None or registry.professional is None:
        return False
    return bool(registry.professional.run_id) and registry.total_count > 0


def export_filename(project: str, month_display: str) -> str:
    canonical = normalize_month_key(str(month_display or "").replace(" ", "-")) or "month"
    token = re.sub(r"[^\w\-]+", "_", str(project or "").strip(), flags=re.UNICODE).strip("_")
    return f"{token or 'project'}_{canonical}_ПРЕДВАРИТЕЛЬНОЕ_МЕСЯЧНОЕ_НАМЕРЕНИЕ_БЕЗ_ДОПУСКА.xlsx"


def scope_lines(context: Optional[WorkingContext]) -> tuple[str, ...]:
    if context is None:
        return (
            "Очередь: Все очереди",
            "Титулы: Все титулы",
            "Дисциплина: Все дисциплины",
            "Системы: Все системы",
            "Пакеты: Все пакеты",
        )
    return (
        f"Очередь: {context.queue or 'Все очереди'}",
        f"Титулы: {', '.join(context.titles) if context.titles else 'Все титулы'}",
        f"Дисциплина: {context.discipline or 'Все дисциплины'}",
        f"Системы: {', '.join(context.systems) if context.systems else 'Все системы'}",
        f"Пакеты: {context.packages or 'Все пакеты'}",
    )


def build_pre_admission_workbook(
    registry: ConstructorDecisionRegistryView,
    *,
    project: str,
    month_display: str,
    context: Optional[WorkingContext] = None,
    generated_at: Optional[datetime] = None,
) -> bytes:
    """Build one workbook from an already loaded professional result."""
    if not export_available(registry):
        raise ValueError("professional result is required")
    stamp = generated_at or datetime.now(timezone.utc)
    professional = registry.professional
    attention = tuple(
        item
        for item in registry.records
        if constructor_row_awaits_human_review(item, professional)
    )
    attention_ids = {item.source_identity for item in attention}
    included = _sorted_field_rows(
        item
        for item in registry.records
        if item.outcome == "CANDIDATE" and item.source_identity not in attention_ids
    )
    excluded = _sorted_field_rows(
        item for item in registry.records if str(item.outcome).startswith("EXCLUDED_")
    )
    workbook = Workbook()
    summary = workbook.active
    summary.title = "00_Свод"
    _write_summary(
        summary,
        registry=registry,
        professional=professional,
        included=included,
        project=project,
        month_display=month_display,
        context=context,
        generated_at=stamp,
    )
    disciplines = sorted({item.discipline.strip() for item in included})
    for index, discipline in enumerate(disciplines, start=1):
        rows = tuple(item for item in included if item.discipline.strip() == discipline)
        sheet = workbook.create_sheet(_discipline_sheet_name(index, discipline))
        _write_register(sheet, rows, table_name=f"Included{index:02d}")
    if attention:
        sheet = workbook.create_sheet("98_Требует_решения")
        _write_trace(sheet, attention, _ATTENTION_COLUMNS, table_name="Attention", attention=True)
    if excluded:
        sheet = workbook.create_sheet("99_Исключено")
        _write_trace(sheet, excluded, _EXCLUDED_COLUMNS, table_name="Excluded", attention=False)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_summary(
    sheet: Worksheet,
    *,
    registry: ConstructorDecisionRegistryView,
    professional: ConstructorProfessionalResultView,
    included: tuple[ConstructorRowDecisionView, ...],
    project: str,
    month_display: str,
    context: Optional[WorkingContext],
    generated_at: datetime,
) -> None:
    sheet["A1"] = excel_text(PRE_ADMISSION_TITLE)
    sheet["A1"].font = _TITLE_FONT
    sheet["A1"].fill = _TITLE_FILL
    sheet["A2"] = excel_text(PRE_ADMISSION_NOTICE)
    sheet["A2"].fill = _NOTICE_FILL
    sheet["A2"].alignment = _WRAP
    facts = (
        ("Статус", PRE_ADMISSION_STATUS),
        ("Проект", project),
        ("Месяц", month_display),
        ("Дата/время формирования", generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Run ID", professional.run_id),
        ("Mission ID", professional.mission_id),
        ("Цифровой сотрудник", EMPLOYEE_ROLE_RU),
        ("Admission status", ADMISSION_STATUS),
    )
    row = 4
    for label, value in facts:
        sheet.cell(row, 1, excel_text(label)).font = _HEADER_FONT
        _write_value(sheet.cell(row, 2), value)
        row += 1
    row += 1
    sheet.cell(row, 1, excel_text("Scope")).font = _HEADER_FONT
    row += 1
    for line in scope_lines(context):
        sheet.cell(row, 1, excel_text(line))
        row += 1
    row += 1
    header_row = row
    for column, title in enumerate(_SUMMARY_COLUMNS, start=1):
        cell = sheet.cell(header_row, column, excel_text(title))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    grouped: dict[str, list[ConstructorRowDecisionView]] = {}
    for item in included:
        grouped.setdefault(item.discipline.strip(), []).append(item)
    for offset, discipline in enumerate(sorted(grouped)):
        rows = grouped[discipline]
        volume, unit = _aggregatable_volume(rows)
        values = (
            discipline or "Без указанной дисциплины",
            len(rows),
            volume,
            unit,
            len({item.facility for item in rows if item.facility}),
            len({item.system for item in rows if item.system}),
            len({item.iwp for item in rows if item.iwp}),
        )
        for column, value in enumerate(values, start=1):
            _write_value(sheet.cell(header_row + 1 + offset, column), value)
    last = header_row + max(len(grouped), 1)
    if grouped:
        table = Table(displayName="SummaryByDiscipline", ref=f"A{header_row}:{get_column_letter(len(_SUMMARY_COLUMNS))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "A4"
    sheet.row_dimensions[2].height = 32
    _widths(sheet, (42, 56, 22, 18, 24, 22, 22))
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_register(
    sheet: Worksheet,
    records: tuple[ConstructorRowDecisionView, ...],
    *,
    table_name: str,
) -> None:
    _write_notice(sheet)
    header_row = 4
    for column, title in enumerate(_FIELD_COLUMNS, start=1):
        cell = sheet.cell(header_row, column, excel_text(title))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
    for offset, record in enumerate(records, start=1):
        values = (
            offset,
            record.facility,
            record.system or "",
            record.iwp,
            record.boq_code,
            record.name,
            record.unit,
            _number(record.remaining_qty),
            _number(record.available_qty),
            None,
            None,
            *(None,) * _WEEK_COLUMNS,
            None,
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(header_row + offset, column)
            _write_value(cell, value)
            cell.alignment = _WRAP
    last = header_row + max(len(records), 1)
    if records:
        table = Table(
            displayName=table_name,
            ref=f"A{header_row}:{get_column_letter(len(_FIELD_COLUMNS))}{last}",
        )
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "A5"
    _widths(
        sheet,
        (6, 16, 18, 16, 18, 36, 12, 14, 28, 24, 20, 12, 12, 12, 12, 12, 24),
    )
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.print_title_rows = "1:4"
    sheet.oddHeader.left.text = "ПРЕДВАРИТЕЛЬНОЕ МЕСЯЧНОЕ НАМЕРЕНИЕ · БЕЗ ДОПУСКА"


def _write_trace(
    sheet: Worksheet,
    records: tuple[ConstructorRowDecisionView, ...],
    columns: tuple[str, ...],
    *,
    table_name: str,
    attention: bool,
) -> None:
    _write_notice(sheet)
    header_row = 4
    for column, title in enumerate(columns, start=1):
        cell = sheet.cell(header_row, column, excel_text(title))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
    ordered = _sorted_field_rows(records)
    for offset, record in enumerate(ordered, start=1):
        if attention:
            values = (
                record.facility,
                record.discipline,
                record.system or "",
                record.iwp,
                record.boq_code,
                record.name,
                _number(record.available_qty),
                record.reason_text,
                "Требует решения",
            )
        else:
            values = (
                record.facility,
                record.discipline,
                record.system or "",
                record.iwp,
                record.boq_code,
                record.name,
                constructor_row_outcome_label(record.outcome),
                record.reason_text,
            )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(header_row + offset, column)
            _write_value(cell, value)
            cell.alignment = _WRAP
    last = header_row + max(len(ordered), 1)
    if ordered:
        table = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(len(columns))}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "A5"
    _widths(sheet, (16, 22, 18, 16, 18, 36, 22, 40, 24))


def _write_notice(sheet: Worksheet) -> None:
    sheet["A1"] = excel_text(PRE_ADMISSION_NOTICE)
    sheet["A1"].fill = _NOTICE_FILL
    sheet["A1"].alignment = _WRAP
    sheet["A2"] = excel_text(f"ADMISSION_STATUS = {ADMISSION_STATUS}")
    sheet["A2"].font = _HEADER_FONT
    sheet.row_dimensions[1].height = 32


def _sorted_field_rows(records) -> tuple[ConstructorRowDecisionView, ...]:
    return tuple(
        sorted(
            records,
            key=lambda item: (
                str(item.facility or ""),
                str(item.system or ""),
                str(item.iwp or ""),
                str(item.boq_code or ""),
                int(item.source_order),
            ),
        )
    )


def _aggregatable_volume(
    records: tuple[ConstructorRowDecisionView, ...] | list[ConstructorRowDecisionView],
) -> tuple[object, str]:
    units = {str(item.unit or "").strip() for item in records if str(item.unit or "").strip()}
    numbers = [_number(item.available_qty) for item in records]
    numeric = [item for item in numbers if isinstance(item, float)]
    if len(units) == 1 and len(numeric) == len(records) and records:
        return round(sum(numeric), 6), next(iter(units))
    return "", ""


def _discipline_sheet_name(index: int, discipline: str) -> str:
    cleaned = _SHEET_FORBIDDEN.sub(" ", discipline).strip() or "Без указанной дисциплины"
    return f"{index:02d}_{cleaned}"[:31]


def _number(value: Optional[str]) -> object:
    text = "" if value is None else str(value).strip().replace(" ", "")
    if not text or not _NUMBER_RE.match(text):
        return excel_text(text)
    return float(text.replace(",", "."))


def excel_text(value: object) -> str:
    """Keep source text, but do not let Excel treat it as a formula."""
    if value is None:
        return ""
    text = str(value)
    if text[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _write_value(cell, value: object) -> None:
    if value is None:
        cell.value = None
        return
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
        if isinstance(value, float):
            cell.number_format = "0.###"
        return
    cell.value = excel_text(value)
    cell.alignment = _WRAP


def _widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
