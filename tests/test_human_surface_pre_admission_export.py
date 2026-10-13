"""Temporary pre-admission workbook. Read-only; no product database."""

from __future__ import annotations

import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from agents.control_room.dtos import (
    ConstructorDecisionRegistryView,
    ConstructorProfessionalResultView,
    ConstructorRowDecisionView,
    ProfessionalResultAvailability,
)
from agents.human_surface.mission_draft import WorkingContext
from agents.human_surface.pre_admission_export import (
    ADMISSION_STATUS,
    EXPORT_BUTTON_LABEL,
    PRE_ADMISSION_NOTICE,
    build_pre_admission_workbook,
    excel_text,
    export_available,
    export_filename,
)
from agents.human_surface.workforce import (
    contour_option_label,
    employee_option_label,
)
from agents.human_surface.workforce import CONTOURS


ROOT = Path(__file__).resolve().parents[1]


def _record(
    order: int,
    outcome: str,
    *,
    discipline: str,
    facility: str = "16160-13",
    boq: str = "",
    name: str = "Работа",
    unit: str = "м2",
    system: str = "АСУТП",
    iwp: str = "IWP-13",
) -> ConstructorRowDecisionView:
    return ConstructorRowDecisionView(
        source_identity=f"{order}|{facility}|{boq or outcome}",
        source_order=order,
        boq_code=boq or f"BOQ-{order}",
        name=name,
        unit=unit,
        system=system,
        remaining_qty="2",
        already_planned_qty="0.5",
        available_qty="1.5",
        outcome=outcome,
        reason_code=outcome,
        reason_text="=CMD()" if order == 1 else "Сохранённое основание",
        requires_human=outcome == "CANDIDATE",
        candidate_created=outcome == "CANDIDATE",
        candidate_id="c-1" if outcome == "CANDIDATE" else "",
        facility=facility,
        discipline=discipline,
        iwp=iwp,
    )


def _registry(
    records: tuple[ConstructorRowDecisionView, ...],
    *,
    status: str = "PACKAGE_BUILT",
) -> ConstructorDecisionRegistryView:
    result = ConstructorProfessionalResultView(
        availability=ProfessionalResultAvailability.AVAILABLE,
        run_id="run-export-1",
        mission_id="mission-export-1",
        professional_status=status,
        project_id="PRJ_001_БХК",
        month="октябрь-2026",
        facility_id="",
        discipline="",
        source_candidate_package_id="pkg-1",
        reviewed_package_id="",
        included_count=None,
        excluded_count=None,
        unresolved_count=None,
        confirm_event_id="",
        confirm_decision="",
        confirm_actor_type="",
        confirm_actor_id="",
        path_statuses=("WAITING_FOR_HUMAN_REVIEW",),
        candidates=(),
        package_candidate_count=0,
    )
    included = sum(1 for item in records if item.outcome == "CANDIDATE")
    excluded = sum(1 for item in records if item.outcome.startswith("EXCLUDED_"))
    return ConstructorDecisionRegistryView(
        availability=ProfessionalResultAvailability.AVAILABLE,
        professional=result,
        total_count=len(records),
        outcome_counts=(("CANDIDATE", included), ("EXCLUDED_COMPLETED", excluded)),
        known_systems=("АСУТП",),
        attention_count=0,
        records=records,
    )


class PreAdmissionExportTests(unittest.TestCase):
    def test_labels_live_inside_the_dropdown_helpers(self) -> None:
        labels = tuple(contour_option_label(item) for item in CONTOURS)
        self.assertEqual(
            labels,
            (
                "Месячное производственное обязательство",
                "Исполнение — ещё не подключено",
                "Признание — ещё не подключено",
                "Причинность и эскалация — ещё не подключено",
            ),
        )
        monthly = CONTOURS[0]
        employees = tuple(employee_option_label(item) for item in monthly.employees)
        self.assertEqual(employees[0], "Инженер формирования месячного плана")
        self.assertIn("ещё не подключён", employees[1])
        self.assertIn("ещё не подключён", employees[2])
        page = (ROOT / "pages" / "53_AI_Центр_управления_агентами.py").read_text(encoding="utf-8")
        self.assertNotIn('st.caption(f"{item.title}', page)

    def test_export_requires_the_current_professional_result(self) -> None:
        self.assertFalse(export_available(None))
        self.assertEqual(EXPORT_BUTTON_LABEL, "ВЫГРУЗИТЬ ПРЕДВАРИТЕЛЬНОЕ МЕСЯЧНОЕ НАМЕРЕНИЕ · XLSX")

    def test_workbook_separates_included_and_excluded_without_invented_crews(self) -> None:
        records = (
            _record(1, "CANDIDATE", discipline="Автоматизация", boq="=BOQ", name="+работа"),
            _record(
                4,
                "CANDIDATE",
                discipline="Автоматизация",
                facility="16160-17",
                boq="=BOQ",
                name="Вторая",
                unit="шт",
                system="БС-2",
                iwp="IWP-17",
            ),
            _record(2, "CANDIDATE", discipline="Вентиляция", facility="26160-17", boq="BOQ-2"),
            _record(5, "CANDIDATE", discipline="Отопление", boq="BOQ-HEAT"),
            _record(3, "EXCLUDED_COMPLETED", discipline="КМ", boq="BOQ-3"),
        )
        registry = _registry(records)
        self.assertTrue(export_available(registry))
        context = WorkingContext(project="БХК", planning_month="октябрь 2026")
        payload = build_pre_admission_workbook(
            registry,
            project="БХК",
            month_display="октябрь 2026",
            context=context,
            generated_at=datetime(2026, 10, 12, 12, 0, tzinfo=timezone.utc),
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(
            workbook.sheetnames,
            ["00_Свод", "01_Автоматизация", "02_Вентиляция", "03_Отопление", "99_Исключено"],
        )
        self.assertNotIn("98_Требует_решения", workbook.sheetnames)
        self.assertNotIn("Изоляция", " ".join(workbook.sheetnames))
        summary_sheet = workbook["00_Свод"]
        summary = " ".join(str(cell.value) for row in summary_sheet.iter_rows() for cell in row if cell.value)
        self.assertIn("ПРЕДВАРИТЕЛЬНОЕ МЕСЯЧНОЕ НАМЕРЕНИЕ", summary)
        self.assertIn(PRE_ADMISSION_NOTICE, summary)
        self.assertIn("run-export-1", summary)
        self.assertIn("mission-export-1", summary)
        self.assertIn("БХК", summary)
        self.assertIn("октябрь 2026", summary)
        self.assertIn(ADMISSION_STATUS, summary)
        self.assertIn("Все очереди", summary)
        self.assertIn("Количество рабочих строк", summary)
        self.assertIn("Дата/время формирования", summary)
        mixed = next(row for row in summary_sheet.iter_rows() if row[0].value == "Автоматизация")
        self.assertIn(mixed[2].value, (None, ""))
        ventilation = next(row for row in summary_sheet.iter_rows() if row[0].value == "Вентиляция")
        self.assertEqual(ventilation[2].value, 1.5)
        self.assertEqual(ventilation[3].value, "м2")
        automation = workbook["01_Автоматизация"]
        self.assertIn(PRE_ADMISSION_NOTICE, str(automation["A1"].value))
        self.assertIn(ADMISSION_STATUS, str(automation["A2"].value))
        self.assertEqual(automation.freeze_panes, "A5")
        self.assertTrue(automation.tables)
        headers = [cell.value for cell in automation[4]]
        self.assertEqual(headers[4], "BOQ-код")
        self.assertIn("Ответственный мастер", headers)
        self.assertIn("Звено / бригада", headers)
        self.assertEqual(
            [headers[index] for index in range(11, 16)],
            ["Неделя 1", "Неделя 2", "Неделя 3", "Неделя 4", "Неделя 5"],
        )
        self.assertEqual(automation.cell(5, 2).value, "16160-13")
        self.assertEqual(automation.cell(6, 2).value, "16160-17")
        remainder = automation.cell(5, headers.index("Остаток") + 1).value
        self.assertIsInstance(remainder, (int, float))
        self.assertNotIsInstance(remainder, bool)
        self.assertEqual(remainder, 2)
        for title in ("Ответственный мастер", "Звено / бригада", "Неделя 1", "Неделя 5", "Примечание"):
            self.assertIsNone(automation.cell(5, headers.index(title) + 1).value)
            self.assertIsNone(automation.cell(6, headers.index(title) + 1).value)
        values = [
            [cell.value for cell in row]
            for row in automation.iter_rows(min_row=5, max_row=6)
        ]
        flat = " ".join(str(item) for line in values for item in line)
        self.assertEqual(sum("=BOQ" in str(item) for line in values for item in line), 2)
        self.assertNotIn("Иванов", flat)
        self.assertNotIn("бригада 1", flat)
        self.assertNotIn("Неделя", flat)
        boq_cell = automation.cell(5, headers.index("BOQ-код") + 1)
        self.assertEqual(boq_cell.data_type, "s")
        self.assertFalse(str(boq_cell.value).startswith("="))
        self.assertTrue(str(automation.cell(5, headers.index("Наименование работы") + 1).value).startswith("'"))
        self.assertEqual(excel_text("=1+1"), "'=1+1")
        self.assertEqual(excel_text("+работа"), "'+работа")
        self.assertEqual(excel_text("-1"), "'-1")
        self.assertEqual(excel_text("@cmd"), "'@cmd")
        excluded = workbook["99_Исключено"]
        excluded_text = " ".join(
            str(cell.value) for row in excluded.iter_rows() for cell in row if cell.value
        )
        self.assertIn("BOQ-3", excluded_text)
        self.assertIn("КМ", excluded_text)
        self.assertIn("Не включать", excluded_text)
        self.assertNotIn("EXCLUDED_COMPLETED", excluded_text)
        self.assertNotIn("BOQ-3", flat)
        self.assertNotIn("BOQ-3", " ".join(workbook.sheetnames))
        self.assertEqual(
            export_filename("БХК", "октябрь 2026"),
            "БХК_2026-10_ПРЕДВАРИТЕЛЬНОЕ_МЕСЯЧНОЕ_НАМЕРЕНИЕ_БЕЗ_ДОПУСКА.xlsx",
        )
        source = (ROOT / "agents" / "human_surface" / "pre_admission_export.py").read_text(encoding="utf-8")
        self.assertNotIn("sqlite3", source)
        self.assertNotIn(".commit(", source)

    def test_unresolved_rows_stay_off_production_sheets(self) -> None:
        records = (
            _record(1, "CANDIDATE", discipline="Автоматизация", boq="BOQ-REVIEW"),
            _record(3, "EXCLUDED_COMPLETED", discipline="КМ", boq="BOQ-OUT"),
        )
        registry = _registry(records, status="WAITING_FOR_HUMAN_REVIEW")
        payload = build_pre_admission_workbook(
            registry,
            project="БХК",
            month_display="октябрь 2026",
            generated_at=datetime(2026, 10, 12, 12, 0, tzinfo=timezone.utc),
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ["00_Свод", "98_Требует_решения", "99_Исключено"])
        review = workbook["98_Требует_решения"]
        review_text = " ".join(str(cell.value) for row in review.iter_rows() for cell in row if cell.value)
        self.assertIn("BOQ-REVIEW", review_text)
        self.assertIn("Требует решения", review_text)
        self.assertIn("Почему требуется решение", review_text)
        reason = review.cell(5, 8)
        self.assertEqual(reason.data_type, "s")
        self.assertFalse(str(reason.value).startswith("="))
        excluded = " ".join(
            str(cell.value) for row in workbook["99_Исключено"].iter_rows() for cell in row if cell.value
        )
        self.assertIn("BOQ-OUT", excluded)
        self.assertNotIn("BOQ-REVIEW", excluded)
        self.assertNotIn("BOQ-OUT", review_text)

    def test_inconsistent_quantities_go_to_excluded_not_production(self) -> None:
        records = (
            _record(1, "CANDIDATE", discipline="Вентиляция", boq="BOQ-OK"),
            ConstructorRowDecisionView(
                source_identity="2|16160-13|BOQ-BAD",
                source_order=2,
                boq_code="BOQ-BAD",
                name="Несогласованный остаток",
                unit="м2",
                system="SYS-1",
                remaining_qty="10",
                already_planned_qty="15",
                available_qty="0",
                outcome="EXCLUDED_INCONSISTENT_QUANTITIES",
                reason_code="EXCLUDED_INCONSISTENT_QUANTITIES",
                reason_text=(
                    "already_planned превышает физический остаток; "
                    "строка исключена из Candidate Package без изменения исходных величин."
                ),
                requires_human=True,
                candidate_created=False,
                candidate_id="",
                facility="16160-13",
                discipline="Вентиляция",
                iwp="",
            ),
        )
        registry = _registry(records)
        payload = build_pre_admission_workbook(
            registry,
            project="БХК",
            month_display="октябрь 2026",
            generated_at=datetime(2026, 10, 12, 12, 0, tzinfo=timezone.utc),
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(
            workbook.sheetnames,
            ["00_Свод", "01_Вентиляция", "99_Исключено"],
        )
        self.assertNotIn("98_Требует_решения", workbook.sheetnames)
        production = workbook["01_Вентиляция"]
        production_text = " ".join(
            str(cell.value) for row in production.iter_rows() for cell in row if cell.value
        )
        self.assertIn("BOQ-OK", production_text)
        self.assertNotIn("BOQ-BAD", production_text)
        excluded = workbook["99_Исключено"]
        excluded_text = " ".join(
            str(cell.value) for row in excluded.iter_rows() for cell in row if cell.value
        )
        self.assertIn("BOQ-BAD", excluded_text)
        self.assertIn("already_planned превышает физический остаток", excluded_text)
        self.assertNotIn("EXCLUDED_INCONSISTENT_QUANTITIES", excluded_text)
        # Source quantities stay on the read-model record unchanged.
        bad = next(item for item in registry.records if item.boq_code == "BOQ-BAD")
        self.assertEqual(bad.remaining_qty, "10")
        self.assertEqual(bad.already_planned_qty, "15")
        self.assertEqual(bad.available_qty, "0")


if __name__ == "__main__":
    unittest.main()
